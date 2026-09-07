"""Tests for src/storage/db.py: zero-setup file creation, round-trip
persistence of a full pipeline trace, and the Excel export.

Every test uses a throwaway .db path under pytest's tmp_path -- never the
real config/settings.yaml path -- so the suite never touches (or
requires) data/app.db.
"""

from __future__ import annotations

import json
import socket
import sqlite3
from pathlib import Path

import pytest
from openpyxl import load_workbook

import storage.db as db

DB_MODULE_SOURCE = Path(db.__file__).read_text(encoding="utf-8")


def _result(
    original_text: str = "While in orbit, the star tracker shall track reference stars continuously.",
    needs_human_review: bool = False,
    recommended_score: float = 96.4,
    recommended_index: int = 1,
    vague_term_suggestions: list[dict] | None = None,
    source_location: dict | None = None,
) -> dict:
    """A dict matching exactly what src/pipeline/graph.py's Finalize node
    (run_requirement/run_workbook) produces -- built by hand here so these
    tests don't depend on running the real pipeline/LLM."""
    candidates = [
        {
            "index": 0, "rewritten_text": "candidate zero text", "pattern": "State-driven",
            "llm_confidence": 0.7, "notes": "", "vague_terms": [], "temperature": 0.2, "seed": 0,
            "score": 80.0, "total_rules": 28, "passed_rule_ids": ["R1"], "failed_rules": [],
            "invented_numbers": [], "has_invented_number": False,
        },
        {
            "index": 1, "rewritten_text": original_text, "pattern": "State-driven",
            "llm_confidence": 0.8, "notes": "", "vague_terms": [], "temperature": 0.4, "seed": 1,
            "score": recommended_score, "total_rules": 28, "passed_rule_ids": ["R1"],
            "failed_rules": [{"id": "R7", "title": "Vague Terms", "category": "Accuracy", "reasons": ["x"]}],
            "invented_numbers": [], "has_invented_number": False,
        },
        {
            "index": 2, "rewritten_text": "candidate two text", "pattern": "State-driven",
            "llm_confidence": 0.6, "notes": "", "vague_terms": [], "temperature": 0.6, "seed": 2,
            "score": 70.0, "total_rules": 28, "passed_rule_ids": [], "failed_rules": [],
            "invented_numbers": [], "has_invented_number": False,
        },
    ]
    return {
        "original_text": original_text,
        "source_location": source_location or {"source": "inline"},
        "rule_flags": [
            {"violation_type": "vague_term", "span": "continuously", "start": 61, "end": 73,
             "confidence": 0.75, "reason": "vague temporal word"}
        ],
        "ears_pattern": {"pattern": "State-driven", "confidence": 0.9, "matched_keywords": ["while"], "reason": "r"},
        "candidates": candidates,
        "recommended_index": recommended_index,
        "recommended_text": candidates[recommended_index]["rewritten_text"],
        "recommended_score": recommended_score,
        "vague_term_suggestions": vague_term_suggestions or [],
        "compliance_threshold": 80.0,
        "needs_human_review": needs_human_review,
    }


@pytest.fixture
def db_path(tmp_path) -> Path:
    return tmp_path / "app.db"


@pytest.fixture
def conn(db_path):
    connection = db.connect(db_path)
    yield connection
    connection.close()


# ---------------------------------------------------------------------------
# Zero setup / zero network
# ---------------------------------------------------------------------------


class TestZeroSetup:
    def test_db_file_does_not_exist_until_first_connect(self, db_path):
        assert not db_path.exists()
        conn = db.connect(db_path)
        assert db_path.exists()
        conn.close()

    def test_parent_directory_is_created_automatically(self, tmp_path):
        nested_path = tmp_path / "nested" / "dirs" / "app.db"
        assert not nested_path.parent.exists()
        conn = db.connect(nested_path)
        assert nested_path.exists()
        conn.close()

    def test_connecting_twice_is_idempotent(self, db_path):
        conn1 = db.connect(db_path)
        db.create_run(conn1, file_name="a.xlsx")
        conn1.close()

        conn2 = db.connect(db_path)  # schema already exists -- must not error
        assert len(db.list_runs(conn2)) == 1
        conn2.close()

    def test_resolve_db_path_uses_config_settings_yaml_by_default(self):
        from config import get_settings

        configured = get_settings()["storage"]["sqlite_path"]
        resolved = db.resolve_db_path()
        assert resolved == db._REPO_ROOT / configured

    def test_open_db_context_manager_closes_connection(self, db_path):
        with db.open_db(db_path) as connection:
            db.create_run(connection, file_name="a.xlsx")
        with pytest.raises(sqlite3.ProgrammingError):
            connection.execute("SELECT 1")  # closed connections raise on use


class TestNoNetworkCalls:
    def test_module_imports_no_networking_libraries(self):
        forbidden = ["requests", "httpx", "urllib.request", "aiohttp", "socket.create_connection"]
        for name in forbidden:
            assert name not in DB_MODULE_SOURCE, f"unexpected network-capable import: {name}"

    def test_full_round_trip_never_opens_a_socket(self, db_path, tmp_path, monkeypatch):
        def _blocked(*args, **kwargs):
            raise AssertionError("storage/db.py must never open a network socket")

        monkeypatch.setattr(socket.socket, "connect", _blocked)
        monkeypatch.setattr(socket, "create_connection", _blocked)

        conn = db.connect(db_path)
        run_id = db.save_pipeline_run(conn, file_name="reqs.xlsx", results=[_result(), _result()])
        db.get_requirements_for_run(conn, run_id)
        db.export_run_to_excel(conn, run_id, tmp_path / "out.xlsx")
        conn.close()  # reaching here without the monkeypatched blocker firing is the assertion


# ---------------------------------------------------------------------------
# Round-trip persistence
# ---------------------------------------------------------------------------


class TestRunMetadata:
    def test_create_run_stores_file_metadata(self, conn, tmp_path):
        real_file = tmp_path / "requirements.xlsx"
        real_file.write_bytes(b"pretend xlsx bytes")

        run_id = db.create_run(conn, file_name="requirements.xlsx", file_path=real_file)
        run = db.get_run(conn, run_id)

        assert run["file_name"] == "requirements.xlsx"
        assert run["file_size_bytes"] == len(b"pretend xlsx bytes")
        assert run["sha256"] == __import__("hashlib").sha256(b"pretend xlsx bytes").hexdigest()
        assert run["uploaded_at"]
        assert run["requirement_count"] == 0

    def test_create_run_without_a_real_file_has_no_size_or_hash(self, conn):
        run_id = db.create_run(conn, file_name="typed-directly")
        run = db.get_run(conn, run_id)
        assert run["file_size_bytes"] is None
        assert run["sha256"] is None

    def test_explicit_size_and_hash_are_not_overwritten_by_auto_detection(self, conn, tmp_path):
        real_file = tmp_path / "f.xlsx"
        real_file.write_bytes(b"data")
        run_id = db.create_run(
            conn, file_name="f.xlsx", file_path=real_file, file_size_bytes=999, sha256="deadbeef"
        )
        run = db.get_run(conn, run_id)
        assert run["file_size_bytes"] == 999
        assert run["sha256"] == "deadbeef"

    def test_list_runs_returns_every_run_in_id_order(self, conn):
        id1 = db.create_run(conn, file_name="a.xlsx")
        id2 = db.create_run(conn, file_name="b.xlsx")
        runs = db.list_runs(conn)
        assert [r["id"] for r in runs] == [id1, id2]

    def test_get_run_returns_none_for_unknown_id(self, conn):
        assert db.get_run(conn, 999) is None

    def test_new_run_starts_pending_with_no_error(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        run = db.get_run(conn, run_id)
        assert run["status"] == "pending"
        assert run["error_message"] is None
        assert run["total_requirements"] is None

    def test_update_run_status_transitions_and_records_error(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        db.update_run_status(conn, run_id, "processing")
        assert db.get_run(conn, run_id)["status"] == "processing"

        db.update_run_status(conn, run_id, "failed", error_message="RuntimeError: boom")
        run = db.get_run(conn, run_id)
        assert run["status"] == "failed"
        assert run["error_message"] == "RuntimeError: boom"

    def test_update_run_status_rejects_unknown_status(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        with pytest.raises(ValueError, match="not-a-real-status"):
            db.update_run_status(conn, run_id, "not-a-real-status")

    def test_set_run_total_requirements(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        db.set_run_total_requirements(conn, run_id, 12)
        assert db.get_run(conn, run_id)["total_requirements"] == 12

    def test_save_pipeline_run_marks_status_completed_with_total(self, conn):
        run_id = db.save_pipeline_run(conn, file_name="a.xlsx", results=[_result(), _result()])
        run = db.get_run(conn, run_id)
        assert run["status"] == "completed"
        assert run["total_requirements"] == 2
        assert run["requirement_count"] == 2


class TestRequirementTrace:
    def test_save_and_retrieve_round_trips_every_field(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        result = _result(needs_human_review=True, vague_term_suggestions=[
            {"term": "continuously", "suggestion": "specify a sampling rate"}
        ])

        req_id = db.save_requirement(conn, run_id, sequence_in_run=0, result=result)
        fetched = db.get_requirement(conn, req_id)

        assert fetched["original_text"] == result["original_text"]
        assert fetched["source_location"] == result["source_location"]
        assert fetched["ears_pattern"] == result["ears_pattern"]
        assert fetched["rule_flags"] == result["rule_flags"]
        assert fetched["candidates"] == result["candidates"]
        assert fetched["recommended_index"] == result["recommended_index"]
        assert fetched["recommended_text"] == result["recommended_text"]
        assert fetched["recommended_score"] == result["recommended_score"]
        assert fetched["vague_term_suggestions"] == result["vague_term_suggestions"]
        assert fetched["compliance_threshold"] == result["compliance_threshold"]
        assert fetched["needs_human_review"] is True
        assert isinstance(fetched["needs_human_review"], bool)

    def test_needs_human_review_false_round_trips_as_false_not_null(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        req_id = db.save_requirement(conn, run_id, 0, _result(needs_human_review=False))
        fetched = db.get_requirement(conn, req_id)
        assert fetched["needs_human_review"] is False

    def test_save_pipeline_run_saves_every_result_and_updates_count(self, conn):
        results = [_result(original_text=f"Requirement {i} shall do a thing continuously.") for i in range(4)]
        run_id = db.save_pipeline_run(conn, file_name="batch.xlsx", results=results)

        run = db.get_run(conn, run_id)
        assert run["requirement_count"] == 4

        fetched = db.get_requirements_for_run(conn, run_id)
        assert len(fetched) == 4
        assert [r["sequence_in_run"] for r in fetched] == [0, 1, 2, 3]
        assert [r["original_text"] for r in fetched] == [r["original_text"] for r in results]

    def test_requirements_are_scoped_to_their_own_run(self, conn):
        run1 = db.save_pipeline_run(conn, file_name="a.xlsx", results=[_result(original_text="A shall x.")])
        run2 = db.save_pipeline_run(conn, file_name="b.xlsx", results=[_result(original_text="B shall y.")])

        assert [r["original_text"] for r in db.get_requirements_for_run(conn, run1)] == ["A shall x."]
        assert [r["original_text"] for r in db.get_requirements_for_run(conn, run2)] == ["B shall y."]

    def test_get_requirement_returns_none_for_unknown_id(self, conn):
        assert db.get_requirement(conn, 999) is None

    def test_nested_json_fields_survive_a_process_restart(self, db_path):
        """Closes and reopens the connection (simulating a real restart)
        to make sure nothing was relying on in-memory state."""
        conn1 = db.connect(db_path)
        run_id = db.save_pipeline_run(conn1, file_name="a.xlsx", results=[_result()])
        conn1.close()

        conn2 = db.connect(db_path)
        fetched = db.get_requirements_for_run(conn2, run_id)
        assert len(fetched) == 1
        assert fetched[0]["candidates"][0]["rewritten_text"] == "candidate zero text"
        conn2.close()


# ---------------------------------------------------------------------------
# Requirement lifecycle: analyze -> (generate | edit)
# ---------------------------------------------------------------------------


def _analyzed_result(original_text: str = "TBD, needs work.") -> dict:
    """Shaped like src/pipeline/graph.py's analyze_requirement() output --
    candidates=[], recommended_index=-1, status="analyzed"."""
    return {
        "original_text": original_text,
        "source_location": {"source": "inline"},
        "rule_flags": [],
        "ears_pattern": {"pattern": "Unclear", "confidence": 0.0, "reason": "no shall clause"},
        "candidates": [],
        "recommended_index": -1,
        "recommended_text": original_text,
        "recommended_score": 42.0,
        "vague_term_suggestions": [],
        "compliance_threshold": 80.0,
        "needs_human_review": True,
        "violations": [{"id": "R7", "title": "Vague Terms", "category": "Accuracy", "reasons": ["x"]}],
        "status": "analyzed",
    }


class TestRequirementLifecycle:
    def test_analyzed_row_has_no_candidates_and_carries_violations(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        req_id = db.save_requirement(conn, run_id, 0, _analyzed_result())
        fetched = db.get_requirement(conn, req_id)

        assert fetched["status"] == "analyzed"
        assert fetched["candidates"] == []
        assert fetched["recommended_index"] == -1
        assert len(fetched["violations"]) == 1
        assert fetched["violations"][0]["id"] == "R7"

    def test_update_requirement_status_changes_only_status_and_error(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        req_id = db.save_requirement(conn, run_id, 0, _analyzed_result())

        db.update_requirement_status(conn, req_id, "generating")
        fetched = db.get_requirement(conn, req_id)
        assert fetched["status"] == "generating"
        assert fetched["original_text"] == "TBD, needs work."  # untouched

        db.update_requirement_status(conn, req_id, "failed", error_message="AI_SERVICE_UNAVAILABLE")
        fetched = db.get_requirement(conn, req_id)
        assert fetched["status"] == "failed"
        assert fetched["error_message"] == "AI_SERVICE_UNAVAILABLE"

    def test_update_requirement_status_rejects_unknown_status(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        req_id = db.save_requirement(conn, run_id, 0, _analyzed_result())
        with pytest.raises(ValueError, match="Unknown requirement status"):
            db.update_requirement_status(conn, req_id, "bogus")

    def test_apply_generation_result_overwrites_candidates_and_sets_generated(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        req_id = db.save_requirement(conn, run_id, 0, _analyzed_result())

        generated = _result(needs_human_review=False)
        db.apply_generation_result(conn, req_id, generated)
        fetched = db.get_requirement(conn, req_id)

        assert fetched["status"] == "generated"
        assert fetched["candidates"] == generated["candidates"]
        assert fetched["recommended_text"] == generated["recommended_text"]
        assert fetched["needs_human_review"] is False
        # original_text is NOT overwritten by generation -- it's the
        # source-of-truth Excel cell, generate_requirement() never touches it.
        assert fetched["original_text"] == "TBD, needs work."

    def test_apply_manual_edit_rescopes_as_edited_no_llm_involved(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        req_id = db.save_requirement(conn, run_id, 0, _analyzed_result())

        edited_text = "When the fuel level drops below reserve, the system shall alert the pilot."
        updated = db.apply_manual_edit(conn, req_id, edited_text)

        assert updated["status"] == "edited"
        assert updated["recommended_text"] == edited_text
        assert updated["recommended_index"] == -1
        assert updated["needs_human_review"] is False
        assert updated["recommended_score"] > 42.0  # real re-score, not left stale
        assert updated["candidates"] == []  # no LLM candidates were ever generated

    def test_status_defaults_to_generated_for_old_style_results_without_it(self, conn):
        # A result dict shaped like the pre-split run_requirement() output
        # (no "status"/"violations" keys at all) must still save cleanly.
        run_id = db.create_run(conn, file_name="a.xlsx")
        old_style = _result()
        assert "status" not in old_style and "violations" not in old_style
        req_id = db.save_requirement(conn, run_id, 0, old_style)
        fetched = db.get_requirement(conn, req_id)
        assert fetched["status"] == "generated"
        assert fetched["violations"] == []


class TestSchemaMigration:
    def test_pre_existing_database_gets_new_columns_added(self, db_path):
        """Simulates a database file created before status/violations_json/
        error_message existed -- db.connect() must add them without losing
        any existing row, and must correctly infer "generated" (not the
        column's own default "analyzed") for a row that already has real
        LLM candidates from the old one-shot pipeline."""
        # Build a pre-migration database by hand: the old schema, one row
        # with real candidates already in it (as the old full pipeline
        # would have left it).
        raw = sqlite3.connect(str(db_path))
        raw.executescript(
            """
            CREATE TABLE runs (
                id INTEGER PRIMARY KEY AUTOINCREMENT, file_name TEXT NOT NULL,
                file_path TEXT, file_size_bytes INTEGER, sha256 TEXT,
                uploaded_at TEXT NOT NULL, requirement_count INTEGER NOT NULL DEFAULT 0,
                total_requirements INTEGER, status TEXT NOT NULL DEFAULT 'pending',
                error_message TEXT
            );
            CREATE TABLE requirements (
                id INTEGER PRIMARY KEY AUTOINCREMENT, run_id INTEGER NOT NULL,
                sequence_in_run INTEGER NOT NULL, original_text TEXT NOT NULL,
                source_location_json TEXT NOT NULL, ears_pattern_json TEXT NOT NULL,
                rule_flags_json TEXT NOT NULL, candidates_json TEXT NOT NULL,
                recommended_index INTEGER NOT NULL, recommended_text TEXT NOT NULL,
                recommended_score REAL NOT NULL, vague_term_suggestions_json TEXT NOT NULL,
                compliance_threshold REAL NOT NULL, needs_human_review INTEGER NOT NULL,
                created_at TEXT NOT NULL
            );
            INSERT INTO runs (file_name, uploaded_at, requirement_count, status)
                VALUES ('legacy.xlsx', '2025-01-01T00:00:00', 1, 'completed');
            INSERT INTO requirements (
                run_id, sequence_in_run, original_text, source_location_json,
                ears_pattern_json, rule_flags_json, candidates_json,
                recommended_index, recommended_text, recommended_score,
                vague_term_suggestions_json, compliance_threshold,
                needs_human_review, created_at
            ) VALUES (
                1, 0, 'The gateway shall log connections.', '{}', '{}', '[]',
                '[{"index": 0, "rewritten_text": "x", "score": 90.0}]',
                0, 'x', 90.0, '[]', 80.0, 0, '2025-01-01T00:00:00'
            );
            """
        )
        raw.commit()
        raw.close()

        conn = db.connect(db_path)
        fetched = db.get_requirements_for_run(conn, 1)
        assert len(fetched) == 1
        assert fetched[0]["original_text"] == "The gateway shall log connections."
        # The whole point: NOT stuck at the ADD COLUMN default "analyzed"
        # despite already having a real candidate and recommended_index.
        assert fetched[0]["status"] == "generated"
        assert fetched[0]["violations"] == []
        # accurate_score_status/accurate_score/accurate_violations/
        # accurate_score_error_message are newer still than status/
        # violations_json -- same migration guard must backfill them too.
        assert fetched[0]["accurate_score_status"] == "not_computed"
        assert fetched[0]["accurate_score"] is None
        assert fetched[0]["accurate_violations"] == []
        assert fetched[0]["accurate_score_error_message"] is None
        # The legacy `runs` table above predates consistency_analyzed_at/
        # consistency_last_error/consistency_contradiction_check_skipped
        # too -- same ALTER-TABLE migration guard must add those without
        # erroring or losing the row.
        run = db.get_run(conn, 1)
        assert run["consistency_analyzed_at"] is None
        assert run["consistency_last_error"] is None
        assert not run["consistency_contradiction_check_skipped"]
        conn.close()


class TestAccurateScoreStorage:
    """src/storage/db.py's start_accurate_score/save_accurate_score/
    fail_accurate_score -- the on-demand 42-rule score's persistence, kept
    entirely separate from the always-computed 28-rule
    recommended_score/violations."""

    def test_defaults_to_not_computed_for_a_fresh_row(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        req_id = db.save_requirement(conn, run_id, 0, _analyzed_result())
        fetched = db.get_requirement(conn, req_id)
        assert fetched["accurate_score_status"] == "not_computed"
        assert fetched["accurate_score"] is None
        assert fetched["accurate_violations"] == []

    def test_start_accurate_score_marks_computing(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        req_id = db.save_requirement(conn, run_id, 0, _analyzed_result())
        db.start_accurate_score(conn, req_id)
        assert db.get_requirement(conn, req_id)["accurate_score_status"] == "computing"

    def test_save_accurate_score_stores_score_and_violations(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        req_id = db.save_requirement(conn, run_id, 0, _analyzed_result())
        db.start_accurate_score(conn, req_id)
        violations = [{"id": "R3", "title": "Appropriate Subject-Verb", "category": "Accuracy", "reasons": ["x"]}]

        updated = db.save_accurate_score(conn, req_id, 95.2, violations)

        assert updated["accurate_score_status"] == "done"
        assert updated["accurate_score"] == 95.2
        assert updated["accurate_violations"] == violations
        assert updated["accurate_score_error_message"] is None
        # the original 28-rule score is completely untouched
        assert updated["recommended_score"] == _analyzed_result()["recommended_score"]

    def test_fail_accurate_score_records_status_and_error(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        req_id = db.save_requirement(conn, run_id, 0, _analyzed_result())
        db.start_accurate_score(conn, req_id)

        db.fail_accurate_score(conn, req_id, "AI_SERVICE_UNAVAILABLE")

        fetched = db.get_requirement(conn, req_id)
        assert fetched["accurate_score_status"] == "failed"
        assert fetched["accurate_score_error_message"] == "AI_SERVICE_UNAVAILABLE"

    def test_start_accurate_score_clears_a_previous_error(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        req_id = db.save_requirement(conn, run_id, 0, _analyzed_result())
        db.fail_accurate_score(conn, req_id, "AI_SERVICE_UNAVAILABLE")

        db.start_accurate_score(conn, req_id)

        fetched = db.get_requirement(conn, req_id)
        assert fetched["accurate_score_status"] == "computing"
        assert fetched["accurate_score_error_message"] is None


class TestConsistencyAnalysisOutcome:
    """db.record_consistency_analysis_outcome() -- lets a caller (src/ui/
    api.py) tell "consistency analysis never ran for this run" apart from
    "it ran and found nothing" apart from "it ran and crashed", which an
    empty requirement_relationships table alone can't distinguish."""

    def test_fresh_run_has_null_timestamps(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        run = db.get_run(conn, run_id)
        assert run["consistency_analyzed_at"] is None
        assert run["consistency_last_error"] is None

    def test_successful_outcome_sets_timestamp_with_no_error(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        db.record_consistency_analysis_outcome(conn, run_id, error=None)
        run = db.get_run(conn, run_id)
        assert run["consistency_analyzed_at"] is not None
        assert run["consistency_last_error"] is None

    def test_failed_outcome_sets_both_timestamp_and_error(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        db.record_consistency_analysis_outcome(conn, run_id, error="boom")
        run = db.get_run(conn, run_id)
        assert run["consistency_analyzed_at"] is not None
        assert run["consistency_last_error"] == "boom"

    def test_a_later_success_clears_a_previous_error(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        db.record_consistency_analysis_outcome(conn, run_id, error="boom")
        db.record_consistency_analysis_outcome(conn, run_id, error=None)
        run = db.get_run(conn, run_id)
        assert run["consistency_last_error"] is None

    def test_fresh_run_has_no_contradiction_check_skipped_flag(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        run = db.get_run(conn, run_id)
        assert not run["consistency_contradiction_check_skipped"]

    def test_records_contradiction_check_skipped(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        db.record_consistency_analysis_outcome(conn, run_id, error=None, contradiction_check_skipped=True)
        run = db.get_run(conn, run_id)
        assert run["consistency_contradiction_check_skipped"] == 1

    def test_a_later_run_with_llm_reachable_clears_the_skipped_flag(self, conn):
        run_id = db.create_run(conn, file_name="a.xlsx")
        db.record_consistency_analysis_outcome(conn, run_id, error=None, contradiction_check_skipped=True)
        db.record_consistency_analysis_outcome(conn, run_id, error=None, contradiction_check_skipped=False)
        run = db.get_run(conn, run_id)
        assert not run["consistency_contradiction_check_skipped"]


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------


class TestExcelExport:
    def test_raises_for_unknown_run_id(self, conn, tmp_path):
        with pytest.raises(ValueError, match="999"):
            db.export_run_to_excel(conn, 999, tmp_path / "out.xlsx")

    def test_creates_file_with_one_row_per_requirement_plus_header(self, conn, tmp_path):
        run_id = db.save_pipeline_run(
            conn, file_name="a.xlsx", results=[_result(), _result(), _result()]
        )
        out_path = db.export_run_to_excel(conn, run_id, tmp_path / "out.xlsx")

        assert out_path.exists()
        wb = load_workbook(out_path)
        ws = wb.active
        assert ws.max_row == 4  # header + 3 requirements

    def test_header_row_matches_spec(self, conn, tmp_path):
        run_id = db.save_pipeline_run(conn, file_name="a.xlsx", results=[_result()])
        out_path = db.export_run_to_excel(conn, run_id, tmp_path / "out.xlsx")
        ws = load_workbook(out_path).active
        headers = [cell.value for cell in ws[1]]
        assert headers == [
            "#", "Source Location", "Original Requirement", "EARS Pattern", "Status",
            "Recommended Requirement", "Compliance Score (28 automatable rules)",
            "Accurate Score (42 rules, incl. AI-judged)",
            "Alternate 1", "Alternate 2", "Needs Review", "Suggestions", "Violations",
        ]

    def test_recommended_and_alternate_columns_are_populated_correctly(self, conn, tmp_path):
        run_id = db.save_pipeline_run(conn, file_name="a.xlsx", results=[_result(recommended_index=1)])
        out_path = db.export_run_to_excel(conn, run_id, tmp_path / "out.xlsx")
        ws = load_workbook(out_path).active

        recommended_cell = ws.cell(row=2, column=6).value
        alt1_cell = ws.cell(row=2, column=9).value
        alt2_cell = ws.cell(row=2, column=10).value

        assert "While in orbit" in recommended_cell  # candidate index 1's text
        # alternates are the two non-recommended candidates, with their scores shown
        assert "candidate zero text" in alt1_cell and "80.0" in alt1_cell
        assert "candidate two text" in alt2_cell and "70.0" in alt2_cell
        # the recommended text itself must not leak into the alternates
        assert "While in orbit" not in alt1_cell
        assert "While in orbit" not in alt2_cell

    def test_status_column_shows_a_human_readable_label(self, conn, tmp_path):
        run_id = db.save_pipeline_run(conn, file_name="a.xlsx", results=[_result()])
        out_path = db.export_run_to_excel(conn, run_id, tmp_path / "out.xlsx")
        ws = load_workbook(out_path).active
        # _result() has no "status" key -> save_requirement() defaults it to
        # "generated" (old-style one-shot pipeline result).
        assert ws.cell(row=2, column=5).value == db._STATUS_LABELS["generated"]

    def test_needs_review_column_and_row_color_for_flagged_requirement(self, conn, tmp_path):
        run_id = db.save_pipeline_run(conn, file_name="a.xlsx", results=[_result(needs_human_review=True)])
        out_path = db.export_run_to_excel(conn, run_id, tmp_path / "out.xlsx")
        ws = load_workbook(out_path).active

        assert ws.cell(row=2, column=11).value == "Yes"
        fill_rgb = ws.cell(row=2, column=11).fill.start_color.rgb
        assert fill_rgb.endswith(db._NEEDS_REVIEW_FILL)

    def test_needs_review_column_and_row_color_for_clean_requirement(self, conn, tmp_path):
        run_id = db.save_pipeline_run(conn, file_name="a.xlsx", results=[_result(needs_human_review=False)])
        out_path = db.export_run_to_excel(conn, run_id, tmp_path / "out.xlsx")
        ws = load_workbook(out_path).active

        assert ws.cell(row=2, column=11).value == "No"
        fill_rgb = ws.cell(row=2, column=11).fill.start_color.rgb
        assert fill_rgb.endswith(db._OK_FILL)

    def test_analyzed_but_not_yet_generated_requirement_is_red(self, conn, tmp_path):
        """The user-visible point of the whole status column: a row that
        was only deterministically analyzed (no Generate/Edit yet) must
        show red, not green, even though nothing "failed" -- it just
        hasn't been resolved by a human yet."""
        run_id = db.create_run(conn, file_name="a.xlsx")
        db.save_requirement(conn, run_id, 0, _analyzed_result())
        out_path = db.export_run_to_excel(conn, run_id, tmp_path / "out.xlsx")
        ws = load_workbook(out_path).active

        assert ws.cell(row=2, column=5).value == db._STATUS_LABELS["analyzed"]
        assert ws.cell(row=2, column=11).value == "Yes"
        fill_rgb = ws.cell(row=2, column=11).fill.start_color.rgb
        assert fill_rgb.endswith(db._NEEDS_REVIEW_FILL)

    def test_edited_requirement_is_green(self, conn, tmp_path):
        run_id = db.create_run(conn, file_name="a.xlsx")
        req_id = db.save_requirement(conn, run_id, 0, _analyzed_result())
        db.apply_manual_edit(conn, req_id, "The gateway shall log every connection attempt.")
        out_path = db.export_run_to_excel(conn, run_id, tmp_path / "out.xlsx")
        ws = load_workbook(out_path).active

        assert ws.cell(row=2, column=5).value == db._STATUS_LABELS["edited"]
        assert ws.cell(row=2, column=11).value == "No"
        fill_rgb = ws.cell(row=2, column=11).fill.start_color.rgb
        assert fill_rgb.endswith(db._OK_FILL)

    def test_entire_row_is_color_highlighted_not_just_one_cell(self, conn, tmp_path):
        run_id = db.save_pipeline_run(conn, file_name="a.xlsx", results=[_result(needs_human_review=True)])
        out_path = db.export_run_to_excel(conn, run_id, tmp_path / "out.xlsx")
        ws = load_workbook(out_path).active

        for col in range(1, len(db._HEADERS) + 1):
            fill_rgb = ws.cell(row=2, column=col).fill.start_color.rgb
            assert fill_rgb.endswith(db._NEEDS_REVIEW_FILL)

    def test_suggestions_column_lists_term_and_suggestion(self, conn, tmp_path):
        run_id = db.save_pipeline_run(
            conn, file_name="a.xlsx",
            results=[_result(vague_term_suggestions=[
                {"term": "continuously", "suggestion": "specify a sampling rate in Hz"}
            ])],
        )
        out_path = db.export_run_to_excel(conn, run_id, tmp_path / "out.xlsx")
        ws = load_workbook(out_path).active
        assert ws.cell(row=2, column=12).value == "continuously: specify a sampling rate in Hz"

    def test_empty_suggestions_render_as_none_placeholder(self, conn, tmp_path):
        run_id = db.save_pipeline_run(conn, file_name="a.xlsx", results=[_result(vague_term_suggestions=[])])
        out_path = db.export_run_to_excel(conn, run_id, tmp_path / "out.xlsx")
        ws = load_workbook(out_path).active
        assert ws.cell(row=2, column=12).value == "(none)"

    def test_violations_column_lists_rule_id_title_and_reason(self, conn, tmp_path):
        run_id = db.create_run(conn, file_name="a.xlsx")
        db.save_requirement(conn, run_id, 0, _analyzed_result())
        out_path = db.export_run_to_excel(conn, run_id, tmp_path / "out.xlsx")
        ws = load_workbook(out_path).active
        assert ws.cell(row=2, column=13).value == "R7 (Vague Terms): x"

    def test_empty_violations_render_as_none_placeholder(self, conn, tmp_path):
        run_id = db.save_pipeline_run(conn, file_name="a.xlsx", results=[_result()])
        out_path = db.export_run_to_excel(conn, run_id, tmp_path / "out.xlsx")
        ws = load_workbook(out_path).active
        assert ws.cell(row=2, column=13).value == "(none)"

    def test_source_location_from_a_real_spreadsheet_cell_is_formatted_as_sheet_and_cell(self, conn, tmp_path):
        location = {
            "file_path": "reqs.xlsx", "sheet_name": "Requirements", "row": 7,
            "column": 3, "column_letter": "C", "is_merged": False, "merge_anchor": None,
        }
        run_id = db.save_pipeline_run(
            conn, file_name="a.xlsx", results=[_result(source_location=location)]
        )
        out_path = db.export_run_to_excel(conn, run_id, tmp_path / "out.xlsx")
        ws = load_workbook(out_path).active
        assert ws.cell(row=2, column=2).value == "Requirements!C7"

    def test_creates_output_parent_directory(self, conn, tmp_path):
        run_id = db.save_pipeline_run(conn, file_name="a.xlsx", results=[_result()])
        nested_out = tmp_path / "exports" / "review.xlsx"
        out_path = db.export_run_to_excel(conn, run_id, nested_out)
        assert out_path.exists()

    def test_exported_file_is_a_valid_readable_workbook(self, conn, tmp_path):
        run_id = db.save_pipeline_run(conn, file_name="a.xlsx", results=[_result()])
        out_path = db.export_run_to_excel(conn, run_id, tmp_path / "out.xlsx")
        # load_workbook already raises on a corrupt file; this just re-confirms
        # a full round trip through openpyxl doesn't lose the sheet.
        wb = load_workbook(out_path)
        assert wb.sheetnames == ["Requirements Review"]
