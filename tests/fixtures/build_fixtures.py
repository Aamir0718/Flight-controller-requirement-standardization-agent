"""Generates the synthetic messy .xlsx fixtures used by
tests/test_parser.py. Run this to (re)materialize the binary fixture
files after changing this script:

    .venv\\Scripts\\python.exe tests\\fixtures\\build_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

FIXTURES_DIR = Path(__file__).parent


def build_messy_multi_sheet() -> None:
    """Two sheets: a real requirements sheet with a merged title banner,
    a blank separator row, a merged note row, a too-short placeholder
    cell, and a second "Legend" sheet that is pure abbreviation
    definitions (should yield ~zero candidates)."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Requirements"
    # NOTE: every row is written via append() (never direct ws["A1"] = ...)
    # so openpyxl's internal row counter and our row comments stay in sync.
    ws.append(["FLIGHT CONTROL SOFTWARE REQUIREMENTS - DRAFT v0.3", None, None, None, None])  # row 1
    ws.merge_cells("A1:E1")
    ws.append([None, None, None, None, None])  # row 2, blank separator (junk row)
    ws.append(["ID", "Category", "Requirement", "Priority", "Notes"])  # row 3
    ws.append([
        "REQ-001", "Nav",
        "The flight control computer shall compute attitude at 50 Hz.",
        "High", "Reviewed",
    ])  # row 4
    ws.append([
        "REQ-002", "Power",
        "While on battery power, the system shall enter low power mode.",
        "Medium", None,
    ])  # row 5
    ws.append([None, None, None, None, None])  # row 6, fully blank junk row
    ws.append(["Note: Section reviewed by QA lead on 2024-01-10", None, None, None, None])  # row 7
    ws.merge_cells("A7:E7")
    ws.append(["REQ-003", "Comms", "TBD", "Low", "pending"])  # row 8, placeholder too short
    ws.append([
        "REQ-004", "Safety",
        "When an overheat condition is detected, the FCC shall shut down the affected channel.",
        "High", None,
    ])  # row 9

    legend = wb.create_sheet("Legend")
    legend.append(["Term", "Meaning"])
    legend.append(["FCC", "Flight Control Computer"])
    legend.append(["TBD", "To Be Determined"])

    wb.save(FIXTURES_DIR / "messy_multi_sheet.xlsx")


def build_header_row_offset() -> None:
    """Header row is not at row 1 (title/blank rows precede it), there is
    a mid-sheet merged section banner, and a second sheet is completely
    empty — both are common in hand-edited DRDO-style sheets."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Section A"
    # NOTE: every row is written via append() (never direct ws["A1"] = ...)
    # so openpyxl's internal row counter and our row comments stay in sync.
    ws.append(["ACME AVIONICS — INTERNAL WORKING COPY", None, None, None])  # row 1
    ws.merge_cells("A1:D1")
    ws.append([None, None, None, None])  # row 2, blank
    ws.append([None, None, None, None])  # row 3, blank
    ws.append(["Req ID", "Requirement Text", "Owner", "Status"])  # row 4: real header
    ws.append([
        "REQ-101",
        "If the primary sensor fails, then the software shall switch to the backup sensor.",
        "J. Rao", "Draft",
    ])  # row 5
    ws.append(["", "", "", ""])  # row 6, blank junk row
    ws.append(["SECTION 2 — NAVIGATION", None, None, None])  # row 7, section banner
    ws.merge_cells("A7:D7")
    ws.append([
        "REQ-102",
        "While in autonomous mode, the navigation unit shall log its position every second.",
        "J. Rao", "Draft",
    ])  # row 8

    empty_sheet = wb.create_sheet("Unused")
    # Deliberately left with zero content to test empty-sheet handling.
    _ = empty_sheet

    wb.save(FIXTURES_DIR / "header_row_offset.xlsx")


def build_edge_cases() -> None:
    """Cell-type and value edge cases: a formula-error sentinel, a
    numeric-only cell, a date cell, a long verb-free paragraph (should be
    skipped despite ample word count), and a genuine requirement mixed
    in among them."""
    import datetime

    wb = Workbook()
    ws = wb.active
    ws.title = "Edge Cases"

    ws.append(["Field", "Value"])
    ws.append(["Formula error", "#REF!"])
    ws.append(["Revision number", 42])
    ws.append(["Effective date", datetime.date(2024, 3, 1)])
    ws.append([
        "Legal footer",
        "This document and its contents remain the confidential property "
        "of the issuing organization and its authorized representatives.",
    ])  # long, word-count passes, but no recognized verb -> skipped
    ws.append([
        "Requirement",
        "If the cabin altitude exceeds 10000 feet, then the system shall "
        "deploy the oxygen masks automatically.",
    ])

    wb.save(FIXTURES_DIR / "edge_cases.xlsx")


def build_corrupt_file() -> None:
    """Not a real workbook at all (plain text saved with an .xlsx
    extension) — exercises the file-open failure path."""
    (FIXTURES_DIR / "corrupt_file.xlsx").write_bytes(
        b"This is not a real xlsx file, just plain bytes.\n"
    )


if __name__ == "__main__":
    build_messy_multi_sheet()
    build_header_row_offset()
    build_edge_cases()
    build_corrupt_file()
    print(f"Fixtures written to {FIXTURES_DIR}")
