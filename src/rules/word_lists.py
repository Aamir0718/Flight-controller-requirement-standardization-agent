"""Loads the plain word/phrase lists src/rules/incose_scorer.py and
src/rules/ears_classifier.py use for their deterministic (no LLM) checks
from data/rules/word_lists.xlsx -- one sheet per list, one term per row
in column A (row 1 is a description, not data).

Why a spreadsheet instead of hardcoding these in Python: every one of
these checks is a plain word/phrase match -- "does this vague term
appear", "is this a known acronym" -- with nothing about it that needs an
LLM or even much code. Keeping the actual words in a spreadsheet means a
domain reviewer (not just a developer) can add or remove a term by
editing a row and saving, no code change or deployment needed.
functools.lru_cache means the file is only actually read once per
process per sheet -- editing it takes effect on the next process start,
same as data/rules/incose_rulebook.json today.

Everything here is still pure, deterministic Python reading a local file
-- no network, no LLM -- exactly like the JSON-backed rule/pattern files
this project already uses (data/rules/incose_rulebook.json,
data/rules/ears_patterns.json).
"""

from __future__ import annotations

import functools
from pathlib import Path

import openpyxl

_WORD_LISTS_PATH = (
    Path(__file__).resolve().parent.parent.parent / "data" / "rules" / "word_lists.xlsx"
)


@functools.lru_cache(maxsize=32)
def load_word_list(sheet_name: str, path: Path | None = None) -> tuple[str, ...]:
    """Returns sheet ``sheet_name``'s column-A terms (row 2 onward; row 1
    is a description, not data) as a tuple, in the order they appear in
    the spreadsheet. Blank cells are skipped. Cached per (sheet_name,
    path) pair.

    Raises KeyError if the workbook has no sheet named ``sheet_name`` --
    listing what sheets it does have, so a typo or a renamed sheet fails
    loudly rather than silently returning an empty/wrong list.
    """
    workbook_path = path or _WORD_LISTS_PATH
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise KeyError(
                f"{workbook_path} has no sheet named {sheet_name!r}; "
                f"available sheets: {workbook.sheetnames}"
            )
        sheet = workbook[sheet_name]
        terms = []
        for row in sheet.iter_rows(min_row=2, min_col=1, max_col=1, values_only=True):
            value = row[0]
            if value is not None and str(value).strip():
                terms.append(str(value).strip())
        return tuple(terms)
    finally:
        workbook.close()
