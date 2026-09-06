"""One-off (and re-runnable) generator for data/rules/word_lists.xlsx --
the editable, non-programmer-facing home for every plain word/phrase list
src/rules/incose_scorer.py and src/rules/ears_classifier.py use for their
deterministic (no LLM) checks: vague terms, escape clauses, banned
abbreviations, known acronyms, weak modal verbs, etc.

Run this again ONLY if you need to regenerate the workbook from scratch
(e.g. lost/corrupted) -- day-to-day list edits (adding a word, an
acronym, ...) should be made directly in data/rules/word_lists.xlsx
itself, one term per row, in the matching sheet. src/rules/word_lists.py
reads that file live; nothing needs re-running after a manual edit.

Usage:
    python scripts/build_word_lists_xlsx.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

REPO_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_PATH = REPO_ROOT / "data" / "rules" / "word_lists.xlsx"

# (sheet_name, description shown in row 1, terms) -- terms start at row 2,
# one per row in column A. Content below is copied byte-for-byte from the
# word lists that were previously hardcoded in incose_scorer.py/
# ears_classifier.py (or, for "acronyms", data/rules/known_abbreviations.json)
# -- this is purely a data-source relocation, not a wording change, so
# every rule keeps scoring exactly the same text exactly the same way.
SHEETS: list[tuple[str, str, list[str]]] = [
    (
        "weak_modals",
        "EARS/INCOSE reserve 'shall' as the sole normative keyword -- these "
        "are the common wrong-modal-verb mistakes (used instead of 'shall').",
        ["will", "must", "should", "may"],
    ),
    (
        "acronyms",
        "INCOSE R37 (Acronyms) -- treated as already defined by domain "
        "convention, so a real aerospace/defense requirement using them "
        "isn't flagged just for not spelling them out inline every time. "
        "Matching is case-insensitive and whole-word. Add project-specific "
        "acronyms here as needed.",
        [
            "GPS", "GNSS", "INS", "IMU", "AHRS", "GPU", "CPU", "MCU", "ECU", "FCC",
            "PWM", "RF", "UART", "SPI", "I2C", "CAN", "USB", "LED", "LCD", "PSI",
            "RPM", "DC", "AC", "ID", "IP", "URL", "RAM", "ROM", "EEPROM", "ADC",
            "DAC", "PID", "FOV", "LOS", "HUD", "IFF", "RADAR", "SONAR", "LIDAR",
            "UAV", "UGV", "USV", "ISR", "C2", "C3I", "GCS", "RTK", "DGPS", "ADS-B",
            "TCAS", "GPWS", "EGPWS", "FMS", "AFCS", "FBW", "APU", "ECS", "FADEC",
            "HSI", "ADI", "VOR", "ILS", "GLONASS", "IRQ", "PCB", "FPGA", "ASIC",
            "SoC", "OS", "RTOS", "API", "SDK", "JSON", "XML", "HTTP", "HTTPS",
            "TCP", "UDP", "SSD", "HDD", "EMI", "EMC", "ESD", "MTBF", "FMEA",
            "DRDO", "ISRO", "DGCA", "MIL-STD", "NATO", "IEEE", "ISO", "INCOSE",
            "EARS",
        ],
    ),
    (
        "vague_terms",
        "INCOSE R7 (Vague Terms) -- unquantifiable qualifiers a requirement "
        "must not use without a measurable criterion.",
        [
            "some", "any", "allowable", "several", "many", "a lot of", "a few",
            "almost always", "very nearly", "nearly", "about", "close to", "almost",
            "approximate", "ancillary", "relevant", "routine", "common", "generic",
            "significant", "flexible", "expandable", "typical", "sufficient",
            "adequate", "appropriate", "efficient", "effective", "proficient",
            "reasonable", "customary",
            "fast", "rapid", "rapidly", "quickly", "swiftly", "promptly",
            "immediately", "instantly", "continuously", "continually",
            "permanently", "negligible", "trivial", "minor", "excessive",
            "extreme", "substantial", "considerable", "normal", "nominal",
            "standard", "comprehensive", "thorough", "detailed", "clear", "crisp",
            "neat", "nicely", "cool", "stuff", "stable", "robust", "accurate",
            "accurately", "precise", "precisely", "safely", "securely", "smoothly",
            "successfully", "properly", "suitably", "correctly", "satisfactorily",
            "user-friendly", "intuitive", "ergonomic", "high accuracy",
            "high precision", "low error", "low latency", "high speed",
            "efficiently", "effectively", "appropriately", "adequately",
            "sufficiently", "reasonably", "significantly", "fully", "completely",
        ],
    ),
    (
        "escape_clauses",
        "INCOSE R8 (Escape Clauses) -- phrases that let a requirement be "
        "satisfied without ever really being met.",
        [
            "so far as is possible", "as little as possible", "where possible",
            "as much as possible", "if it should prove necessary", "if necessary",
            "to the extent necessary", "as appropriate", "as required",
            "to the extent practical", "if practicable",
            "if practical", "if safe", "when safe to do so", "where applicable",
            "when appropriate", "to the extent possible",
        ],
    ),
    (
        "superfluous_infinitives",
        "INCOSE R10 (Superfluous Infinitives).",
        ["to be designed to", "to be able to", "to be capable of", "to enable", "to allow", "be able to", "be capable of"],
    ),
    (
        "combinator_words",
        "INCOSE R19 -- combinator words joining clauses that should be "
        "written as separate requirement statements.",
        ["and", "or", "then", "unless", "but", "as well as", "but also", "however", "whether", "meanwhile", "whereas", "on the other hand", "otherwise"],
    ),
    (
        "purpose_phrases",
        "INCOSE R20 (Purpose Phrases) -- explaining WHY instead of stating WHAT.",
        ["in order to", "so that", "for the purpose of", "the intent of", "the reason for"],
    ),
    (
        "group_noun_references",
        "INCOSE R22 -- referring to a set with a group noun instead of enumerating it.",
        ["the following", "such as", "various", "a variety of", "these parameters", "these items", "several types of"],
    ),
    (
        "personal_pronouns",
        "INCOSE R24 -- personal/indefinite pronouns with an ambiguous antecedent.",
        ["it", "its", "this", "that", "these", "those", "they", "them", "their", "one", "anyone", "someone", "anything", "something", "he", "she", "him", "her"],
    ),
    (
        "unachievable_absolutes",
        "INCOSE R26 -- absolute claims that can never be verified as true.",
        ["100%", "all", "every", "always", "never", "none", "zero defects"],
    ),
    (
        "implied_applicability",
        "INCOSE R27 -- applicability implied rather than stated explicitly.",
        ["typically", "usually", "generally", "normally", "in most cases", "under normal circumstances"],
    ),
    (
        "universal_quantifiers",
        "INCOSE R32 -- use 'each' instead of these for universal quantification.",
        ["all", "any", "both"],
    ),
    (
        "optimization_language",
        "INCOSE R34 -- unbounded performance/optimization language with no measurable target.",
        [
            "optimize", "optimide", "optimizes", "optimized", "optimization",
            "maximize", "maximizes", "maximized", "minimize", "minimizes",
            "minimized", "optimal", "optimum", "improve", "improves", "improved",
            "enhance", "enhances", "enhanced",
        ],
    ),
    (
        "indefinite_temporal_keywords",
        "INCOSE R35 -- indefinite temporal keywords with no measurable timing.",
        [
            "eventually", "until", "before", "after", "as", "once", "earliest",
            "latest", "instantaneous", "simultaneous", "at last",
            "daily", "weekly", "monthly", "periodically", "regularly", "frequently",
            "occasionally", "intermittently", "as soon as possible", "without delay",
        ],
    ),
    (
        "banned_abbreviations",
        "INCOSE R38 (Abbreviations) -- shorthand that should be spelled out.",
        ["approx.", "min.", "max.", "temp.", "qty.", "spec.", "req.", "e.g.", "i.e.", "vs.", "w/o", "w/"],
    ),
    (
        "action_verbs",
        "Recognized action verbs used by the R28 (single-condition-drives-"
        "multiple-actions) compound-requirement check.",
        sorted([
            "activate", "adjust", "aggregate", "alert", "allow", "apply", "arm",
            "assert", "attenuate", "authenticate", "block", "calculate", "capture",
            "clear", "close", "command", "compute", "configure", "correct",
            "deactivate", "decrypt", "deploy", "detect", "disable", "discard",
            "disengage", "display", "double", "drive", "enable", "encrypt",
            "engage", "enter", "execute", "extend", "filter", "flag", "flash",
            "generate", "halt", "highlight", "hold", "identify", "illuminate",
            "indicate", "initialize", "initiate", "isolate", "issue", "limit",
            "log", "maintain", "modulate", "monitor", "notify", "open", "output",
            "overwrite", "parse", "poll", "process", "project", "provide", "pump",
            "read", "record", "reduce", "reject", "release", "render", "report",
            "reset", "resume", "retry", "sample", "save", "scale", "scan", "send",
            "set", "shunt", "shutdown", "sound", "store", "switch", "target",
            "terminate", "track", "transition", "transmit", "trigger", "update",
            "upload", "verify", "write",
        ]),
    ),
]


def build() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)  # replaced by named sheets below

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

    for sheet_name, description, terms in SHEETS:
        sheet = workbook.create_sheet(sheet_name)
        sheet.append([description])
        cell = sheet.cell(row=1, column=1)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.column_dimensions["A"].width = 90
        sheet.row_dimensions[1].height = 45
        for term in terms:
            sheet.append([term])
        sheet.freeze_panes = "A2"

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH} ({len(SHEETS)} sheets, {sum(len(t) for _, _, t in SHEETS)} total terms)")


if __name__ == "__main__":
    build()
